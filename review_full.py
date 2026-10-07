"""
Unified review: data checks + three views per factor, one report for all perils and one per peril.

For each factor:
    1. Raw data                 every value (or fine unrounded-ish bins); N/A and Unknown as their own levels
    2. Suggested bands (Radar)  tidy banding: caps/floors, round-number bands, sparse levels to OTHER.
                                The same for every peril - this is what Radar writes to the Emblem file.
    3. Indicative Emblem groups per peril, merging tidy bands until each has enough claims for that peril.
                                A one-way guide to where Emblem grouping is likely to end up, not a decision.

Every chart: exposure bars stacked by product; burning cost / frequency / severity relativity lines
(per-product burning cost lines can be switched on from the legend). All three charts for a factor use
the same scale: 1.00 = the base level (largest-exposure group, or your choice).

Level names used everywhere (and in the Radar spec):
    N/A      not applicable (product / section) - reserved code from the Formula component
    Unknown  missing where the factor applies, or a junk value
    OTHER    small categorical levels grouped together
"""

import json
import math

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from banding import (C_REF, HtmlReport, _cell, _esc, _fmt, _is_whole, _nice, _slug, _table, _tidy,
                     _weighted_quantiles, candidate_edges, quantile_edges, heat_table, merge_to_credibility, numeric_labels,
                     plot_stability)
from readiness import _colours, _json_rule, _rule_text, reserved_value

C_BC, C_FREQ, C_SEV, C_INK = "#1f1f1d", "#4a3aa7", "#e34948", "#1f1f1d"
SPECIAL = ["Unknown", "N/A"]
NOT_BASE = {"Unknown", "N/A", "OTHER", "(other levels)"}


# =============================================================================== level specs

def _real(df, f, r):
    """Rows where the factor applies and has a value (not N/A, not missing)."""
    s = df[f]
    ok = s.notna()
    if r is not None:
        ok &= ~(s.astype("string") == str(r)).fillna(False) if isinstance(r, str) else ~(s == r).fillna(False)
    return ok


def raw_spec(df, f, cat, r, expo, max_levels=300, bins=100, max_cat_levels=1000):
    """
    No banding unless unavoidable: every categorical level (up to max_cat_levels) and every numeric value
    (up to max_levels distinct values). Only a numeric factor with more distinct values than that is
    binned, into ~bins equal-exposure bins with edges rounded to 2 significant figures.
    """
    ok = _real(df, f, r)
    sub = df.loc[ok]
    if cat:
        g = sub.groupby(sub[f].astype(str))[expo].sum().sort_values(ascending=False)
        top = list(g.index[:max_cat_levels])
        other = "(other levels)"
        mp = {k: (k if k in top else other) for k in g.index}
        return {"type": "categorical", "map": mp, "levels": top + ([other] if len(g) > max_cat_levels else []),
                "other": other, "reserved": r}
    vals = np.sort(sub[f].unique())
    if len(vals) <= max_levels:
        return {"type": "values", "values": [_tidy(v) for v in vals], "reserved": r}
    whole = _is_whole(sub[f])
    q = quantile_edges(sub, f, bins, expo)
    edges = sorted({_sig2(x) for x in q if vals[0] < _sig2(x) <= vals[-1]})
    return {"type": "numeric", "edges": [_tidy(e) for e in edges], "integer": whole,
            "data_min": _tidy(vals[0]) if whole and len(vals) else None, "reserved": r}


def _sig2(x):
    if x == 0:
        return 0.0
    mag = 10 ** (np.floor(np.log10(abs(x))) - 1)
    return float(round(x / mag) * mag)


def _merge_thin(sub, f, edges, expo, claims, min_e, min_n):
    """Merge the thinnest band (exposure or claims below the floor) into its smaller neighbour."""
    edges = sorted(edges)
    while edges:
        b = pd.cut(sub[f], [-np.inf, *edges, np.inf], right=False, labels=False)
        g = sub.groupby(b)[[expo, claims]].sum().reindex(range(len(edges) + 1), fill_value=0)
        bad = (g[expo] < min_e) | (g[claims] < min_n)
        if not bad.any():
            break
        i = int(g[expo][bad].idxmin())
        j = 0 if i == 0 else (i - 1 if i == len(g) - 1 or g[expo][i - 1] <= g[expo][i + 1] else i)
        edges.pop(j)
    return edges


def tidy_spec(df, f, cat, r, expo, claims, s, override=None):
    """
    Cleaning-mode bands for Radar. s: dict(tidy_bins, values_max, min_expo, min_claims, tail).
    override: {"edges": [...]} or {"levels": [...]} to replace the suggestion.
    """
    ok = _real(df, f, r)
    sub = df.loc[ok]
    tot = sub[expo].sum()
    if cat:
        g = sub.groupby(sub[f].astype(str))[[expo, claims]].sum().sort_values(expo, ascending=False)
        if override and "levels" in override:
            keep = [str(x) for x in override["levels"] if str(x) != "OTHER"]
        else:
            keep_m = (g[expo] / tot >= s["min_expo"]) & (g[claims] >= s["min_claims_tidy"])
            keep = list(g.index[keep_m])[:250]
        mp = {k: (k if k in keep else "OTHER") for k in g.index}
        levels = [k for k in g.index if k in keep] + (["OTHER"] if any(v == "OTHER" for v in mp.values()) else [])
        return {"type": "categorical", "map": mp, "levels": levels, "other": "OTHER", "reserved": r,
                "overridden": bool(override)}
    v = sub[f].astype(float)
    w = sub[expo].astype(float)
    vals = np.sort(v.unique())
    whole = _is_whole(v)
    if override and "edges" in override:
        edges = sorted(override["edges"])
    else:
        if whole and len(vals) <= s["values_max"]:
            edges = list(vals[1:])
        elif whole and len(vals) <= 150:
            # whole numbers with a modest range: even steps of 1, 2, 5, 10, 20, 25, 50 ...
            span = vals[-1] - vals[0]
            step = next(st for st in [1, 2, 5, 10, 20, 25, 50, 100, 200, 250, 500, 1000]
                        if span / st <= s["tidy_bins"]) if span <= 1000 * s["tidy_bins"] else None
            if step:
                start = np.floor(vals[0] / step) * step + step
                edges = list(np.arange(start, vals[-1] + 1, step))
            else:
                edges, whole = candidate_edges(sub, f, s["tidy_bins"], expo)
        else:
            edges, whole = candidate_edges(sub, f, s["tidy_bins"], expo)
        if len(vals) > s["values_max"]:                      # cap / floor the thin tails at round numbers
            lo_q, hi_q = _weighted_quantiles(v.to_numpy(), w.to_numpy(), [s["tail"], 1 - s["tail"]])
            lo_c, hi_c = _nice(lo_q), _nice(hi_q)
            if whole:
                lo_c, hi_c = round(lo_c), round(hi_c)
            edges = [e for e in edges if lo_c < e <= hi_c]
            if hi_c < vals[-1]:
                edges.append(hi_c)
            if lo_c > vals[0]:
                edges.append(lo_c)
            edges = sorted({e for e in edges if vals[0] < e <= vals[-1]})
        edges = _merge_thin(sub, f, edges, expo, claims, s["min_expo"] * tot, s["min_claims_tidy"])
    return {"type": "numeric", "edges": [_tidy(e) for e in edges], "integer": whole,
            "data_min": _tidy(vals[0]) if whole and len(vals) else None, "reserved": r, "overridden": bool(override)}


def group_spec(df, f, tidy, expo, claims, cost, min_claims, max_groups):
    """Indicative Emblem grouping for one peril: merge tidy bands until each has min_claims."""
    r = tidy.get("reserved")
    ok = _real(df, f, r)
    sub = df.loc[ok]
    if tidy["type"] == "numeric":
        edges = merge_to_credibility(sub, f, list(tidy["edges"]), min_claims, claims, max_bands=max_groups,
                                     cost=cost, expo=expo)
        return {**tidy, "edges": [_tidy(e) for e in edges]}
    lv = sub[f].astype(str).map(tidy["map"]).fillna("OTHER")
    n = sub.groupby(lv)[claims].sum()
    keep = [l for l in tidy["levels"] if l != "OTHER" and n.get(l, 0) >= min_claims]
    mp = {k: (v if v in keep else "OTHER") for k, v in tidy["map"].items()}
    levels = keep + (["OTHER"] if any(v == "OTHER" for v in mp.values()) else [])
    return {**tidy, "map": mp, "levels": levels}


def level_series(s, spec):
    """Label every row with its level under spec. Returns (labels, ordered level list)."""
    r = spec.get("reserved")
    if spec["type"] == "values":
        na = (s == r).fillna(False) if r is not None and not isinstance(r, str) else pd.Series(False, index=s.index)
        v = pd.to_numeric(s, errors="coerce").where(~na)
        names = {x: _fmt(x) for x in pd.unique(v.dropna())}
        lab = v.map(names).astype(object)
        lab[v.isna().to_numpy()] = "Unknown"
        lab[na.to_numpy()] = "N/A"
        order = list(dict.fromkeys(_fmt(x) for x in spec["values"]))
        present = set(lab.unique())
        return lab, [o for o in order if o in present and o not in SPECIAL] + [o for o in SPECIAL if o in present]
    if spec["type"] == "numeric":
        na = (s == r).fillna(False) if r is not None and not isinstance(r, str) else pd.Series(False, index=s.index)
        v = pd.to_numeric(s, errors="coerce").where(~na)
        names = numeric_labels(spec["edges"], False, spec["integer"], spec.get("data_min"))
        codes = pd.cut(v, [-np.inf, *spec["edges"], np.inf], right=False, labels=False)
        lab = pd.Series([names[int(c)] if pd.notna(c) else "Unknown" for c in codes], index=s.index, dtype=object)
        lab[na] = "N/A"
        order = names
    else:
        st = s.astype("string")
        mapped = st.map(spec["map"])
        lab = mapped.where(mapped.notna(), spec.get("other", "OTHER")).astype(object)
        lab[st.isna().to_numpy()] = "Unknown"
        lab[st.eq("N/A").fillna(False).to_numpy()] = "N/A"
        order = list(spec["levels"])
    present = set(lab.unique())
    return lab, [o for o in order if o in present and o not in SPECIAL] + [o for o in SPECIAL if o in present]


# =============================================================================== tables & charts

def level_table(df, lab, order, expo, claims, cost, prod, products):
    d = pd.DataFrame({"lab": lab, "e": df[expo], "n": df[claims], "c": df[cost], "p": df[prod].astype("string")})
    E, N, C = d.e.sum(), d.n.sum(), d.c.sum()
    g = d.groupby("lab")[["e", "n", "c"]].sum().reindex(order)
    t = pd.DataFrame({"exposure": g.e, "claims": g.n, "cost": g.c})
    t["expo_share"] = t.exposure / E
    t["freq_rel"] = (t.claims / t.exposure) / (N / E)
    t["sev_rel"] = (t.cost / t.claims.where(t.claims > 0)) / (C / N) if N else np.nan
    t["bc_rel"] = (t.cost / t.exposure) / (C / E) if C else np.nan
    pe = d.groupby(["lab", "p"])[["e", "c"]].sum()
    for p in products:
        if p in pe.index.get_level_values(1):
            x = pe.xs(p, level=1).reindex(order)
            t[f"expo_{p}"] = x.e.fillna(0)
            t[f"bc_{p}"] = (x.c / x.e) / (C / E) if C else np.nan
        else:
            t[f"expo_{p}"] = 0.0
            t[f"bc_{p}"] = np.nan
    return t


def other_share(t):
    return float(t.loc["OTHER", "expo_share"]) if "OTHER" in t.index else 0.0


def pick_base_level(t, requested=None):
    if requested is not None and str(requested) in map(str, t.index):
        return str(requested), "your choice"
    cand = t[~t.index.isin(NOT_BASE) & (t.claims > 0)]
    if not len(cand):
        return None, "none available"
    how = "largest exposure" if requested is None else f"'{requested}' not found, so largest exposure"
    return str(cand.exposure.idxmax()), how


def rescale(t, scale):
    t = t.copy()
    for c in [c for c in t.columns if c.endswith("_rel") or c.startswith("bc_")]:
        key = "bc_rel" if c.startswith("bc_") else c
        if scale.get(key) and np.isfinite(scale[key]) and scale[key] > 0:
            t[c] = t[c] / scale[key]
    return t


PLOT_H, PLOT_TOP, ASSUMED_W, TICK_PX = 330, 112, 1100, 11


def fix_plot_area(fig, labels, plot_h=PLOT_H, top=PLOT_TOP):
    """
    Keep the plotting area the same height on every chart: the space the level names need is added
    below it (the figure grows) instead of being taken out of it. Labels turn 40 degrees when they would
    overlap, and vertical when there are very many levels.
    """
    labels = [str(l) for l in labels]
    n, longest = len(labels), max((len(l) for l in labels), default=1)
    char_w = TICK_PX * 0.55
    left = 72
    per_level = (ASSUMED_W - left - 72) / max(n, 1)
    if n > 45:
        angle = -90
    elif longest * char_w > per_level * 0.95:
        angle = -40
    else:
        angle = 0
    sin, cos = abs(math.sin(math.radians(angle))), abs(math.cos(math.radians(angle)))
    bottom = int(longest * char_w * sin + TICK_PX * (cos + 0.6) + 20)
    if angle == -40 and labels:                       # the first label leans left, past the y axis
        left = max(left, int(len(labels[0]) * char_w * cos - per_level / 2 + 12))
    # show every label up to ~80 levels; beyond that every k-th, but always N/A, Unknown and OTHER
    step = max(1, math.ceil(n / 80))
    shown = [l for i, l in enumerate(labels) if i % step == 0 or l in NOT_BASE]
    fig.update_xaxes(tickangle=angle, automargin=False, tickfont=dict(size=TICK_PX),
                     tickmode="array", tickvals=shown, ticktext=shown)
    fig.update_yaxes(automargin=False)
    fig.update_layout(height=plot_h + top + bottom, margin=dict(t=top, b=bottom, l=left, r=72))
    return fig


def plot_levels(t, title, products, colours, min_claims, base=None):
    x = [str(i) for i in t.index]
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    for p in products:
        fig.add_trace(go.Bar(x=x, y=t[f"expo_{p}"], name=p, marker=dict(color=colours[p], opacity=0.38,
                                                                         line=dict(width=0)),
                             hovertemplate=f"{p} exposure %{{y:,.0f}}<extra></extra>"), secondary_y=False)
    thin = (t.claims.fillna(0) < min_claims).to_numpy()
    dense = len(x) > 60                      # hundreds of raw levels: lighter marks so the shape shows
    for c, name, col, width, dash, size in [("bc_rel", "Burning cost", C_BC, 1.5 if dense else 3, "solid", 4 if dense else 9),
                                            ("freq_rel", "Frequency", C_FREQ, 1 if dense else 1.5, "dot", 3 if dense else 6),
                                            ("sev_rel", "Severity", C_SEV, 1 if dense else 1.5, "dash", 3 if dense else 6)]:
        fig.add_trace(go.Scatter(
            x=x, y=t[c], name=name, mode="lines+markers", line=dict(color=col, width=width, dash=dash),
            marker=dict(size=size, color=col, symbol=np.where(thin, "circle-open", "circle")),
            customdata=np.c_[t.claims.fillna(0), t.expo_share.fillna(0) * 100],
            hovertemplate=f"{name} %{{y:.2f}}" + ("<br>Claims %{customdata[0]:,.0f} · exposure %{customdata[1]:.1f}%"
                                                  if c == "bc_rel" else "") + "<extra></extra>"), secondary_y=True)
    for p in products:
        if t[f"bc_{p}"].notna().any():
            fig.add_trace(go.Scatter(x=x, y=t[f"bc_{p}"], name=f"{p} burning cost", mode="lines+markers",
                                     visible="legendonly", line=dict(color=colours[p], width=1.8),
                                     marker=dict(size=6, color=colours[p]),
                                     hovertemplate=f"{p} burning cost %{{y:.2f}}<extra></extra>"), secondary_y=True)
    if base is not None and base in x:
        fig.add_trace(go.Scatter(x=[base], y=[1], mode="markers", name="Base level",
                                 marker=dict(symbol="diamond", size=13, color=C_INK, line=dict(color="white", width=2)),
                                 hoverinfo="skip"), secondary_y=True)
    fig.add_shape(type="line", xref="paper", x0=0, x1=1, yref="y2", y0=1, y1=1,
                  line=dict(color=C_REF, width=1, dash="dot"))
    fig.update_layout(title=dict(text=title, y=0.98, yanchor="top"), template="plotly_white", hovermode="x unified",
                      barmode="stack", bargap=0.12 if len(x) <= 60 else 0.02,
                      legend=dict(orientation="h", x=0, y=1.02, yanchor="bottom", traceorder="normal"))
    fig.update_xaxes(type="category")
    fig.update_yaxes(title_text="Exposure", secondary_y=False, showgrid=False, tickformat=",.0f")
    fig.update_yaxes(title_text="Relativity (1 = base)" if base else "Relativity (1 = average)", secondary_y=True,
                     rangemode="tozero", tickmode="auto", nticks=6, tickformat=".1f", showgrid=True)
    return fix_plot_area(fig, x)


_TFMT = {"exposure": ",.0f", "expo_share": ".1%", "claims": ",.0f", "freq_rel": ".2f", "sev_rel": ".2f", "bc_rel": ".2f"}
_THDR = {"exposure": "Exposure", "expo_share": "Exposure %", "claims": "Claims", "freq_rel": "Frequency rel.",
         "sev_rel": "Severity rel.", "bc_rel": "Burning cost rel."}


def small_table(t, min_claims, base=None):
    tt = t[list(_TFMT)].copy()
    if base is not None:
        tt.index = [f"{i}  (base)" if str(i) == base else i for i in tt.index]
    return _table(tt, _TFMT, _THDR, ["thin" if n < min_claims else "" for n in t.claims.fillna(0)], index_name="Level")


# =============================================================================== per-factor build

def spec_line(f, spec):
    if spec["type"] == "numeric":
        return f'"{f}": {json.dumps({"edges": spec["edges"]})},'
    return f'"{f}": {json.dumps({"levels": [l for l in spec["levels"] if l != "OTHER"]})},'


def factor_views(df, f, cat, tidy, raw, peril, cfg, s):
    """Tables, base and figures for one factor and one peril ('ALL' = all perils combined)."""
    expo, prod, products = cfg["EXPO"], cfg["PRODUCT"], cfg["PRODUCTS"]
    cc, kc = ("_n_all", "_c_all") if peril == "ALL" else cfg["PERIL_COLS"][peril]
    min_c = s["min_claims_group"].get(peril, s["min_claims_group"]["default"])
    grp = group_spec(df, f, tidy, expo, cc, kc, min_c, s["max_groups"])
    tabs = {}
    for k, sp in (("raw", raw), ("tidy", tidy), ("group", grp)):
        lab, order = level_series(df[f], sp)
        tabs[k] = (level_table(df, lab, order, expo, cc, kc, prod, products), lab)
    req = cfg.get("BASE_LEVELS", {}).get(f)
    tidy_base, _ = pick_base_level(tabs["tidy"][0], req) if req is not None and str(req) in map(str, tabs["tidy"][0].index) \
        else (None, None)
    if tidy_base is not None:                       # chosen base is a tidy level: base group = the group containing it
        base = str(pd.crosstab(tabs["tidy"][1], tabs["group"][1]).loc[tidy_base].idxmax())
        how = "your choice"
    else:
        base, how = pick_base_level(tabs["group"][0], req)
        if base is not None:                        # Radar base: largest tidy level inside the base group
            inside = tabs["tidy"][1][tabs["group"][1] == base]
            ex = df.loc[inside.index, expo].groupby(inside).sum()
            ex = ex[~ex.index.isin(NOT_BASE)]
            tidy_base = str(ex.idxmax()) if len(ex) else None
    if base is not None:
        row = tabs["group"][0].loc[base]
        scale = {"bc_rel": row.bc_rel, "freq_rel": row.freq_rel, "sev_rel": row.sev_rel}
        tabs = {k: (rescale(t, scale), lab) for k, (t, lab) in tabs.items()}
    gt = tabs["group"][0]
    cred = gt[(gt.claims >= min_c) & ~gt.index.isin(SPECIAL)]
    spread = (cred.bc_rel.max() / cred.bc_rel.min()) if len(cred) > 1 and cred.bc_rel.min() > 0 else np.nan
    return dict(group=grp, tabs=tabs, base=base, base_how=how, tidy_base=tidy_base, min_c=min_c, spread=spread)


def add_factor_section(rep, df, f, cat, chk, eng_log, raw, tidy, views, peril, cfg, colours, s, formula=None):
    products = cfg["PRODUCTS"]
    tabs = views["tabs"]
    lvl_word = {"action": "Action", "check": "Check", "info": "Info"}
    n_tidy = len([i for i in tabs["tidy"][0].index])
    tiles = [("Type", chk["type"]), ("Distinct", f"{chk['distinct']:,}"), ("Missing (raw)", f"{chk['missing_raw']:.1%}"),
             ("N/A", f"{chk['na_share']:.1%}"), ("Unknown", f"{chk['unknown']:.1%}"),
             ("Modelling bands", f"{n_tidy}"), ("Indicative groups", f"{len(tabs['group'][0])}"),
             ("Base level", views["base"] or "average")]
    intro = ("<div class='stats'>" + "".join(f"<div class='stat'><span>{k}</span><b>{_esc(v)}</b></div>"
                                              for k, v in tiles) + "</div>")
    flags = "".join(f"<li><b>{lvl_word[lv]}:</b> {_esc(t)}</li>" for lv, t in chk["flags"])
    if n_tidy > 255:
        flags += "<li><b>Action:</b> more than 255 modelling levels — Emblem's limit. Group further.</li>"
    oth = other_share(tabs["tidy"][0])
    if oth >= 0.2:
        flags += (f"<li><b>Action:</b> OTHER holds {oth:.0%} of exposure in the suggested bands — too much detail is "
                  f"lost. Group the small levels to a parent first (e.g. trade code → trade group) with MAPPINGS, or "
                  f"lower min_expo / min_claims_tidy.</li>")
    intro += f"<ul class='checks'>{flags or '<li>No data issues found.</li>'}</ul>"
    if eng_log.get(f):
        intro += "<h3>Engineering applied</h3><ul class='checks'>" + "".join(
            f"<li>{_esc(t)}</li>" for t in eng_log[f]) + "</ul>"
    for p in chk.get("proposals", []):
        intro += ("<p class='note'>To apply the proposed recode, add this line to APPLIES_TO and run again:</p>"
                  f"<pre>{_esc(chr(34) + f + chr(34) + ': ' + _json_rule(p['rule']) + ',')}</pre>")
    if formula:
        intro += f"<p class='note'>Radar Formula component (before Banding):</p><pre>{_esc(formula)}</pre>"
    if tidy.get("reserved") is not None:
        intro += ("<p class='note'>In Emblem, fix the N/A level's parameter at 0 (relativity 1.00) with Fixing Factors, "
                  "so the product / section relativity carries those policies.</p>")

    scope = "all perils" if peril == "ALL" else peril
    raw_t, tidy_t, grp_t = tabs["raw"][0], tabs["tidy"][0], tabs["group"][0]
    base = views["base"]
    thin_raw = s["min_claims_tidy"]
    b1 = (f"{f} — raw data ({len(raw_t)} levels, {scope})",
          plot_levels(raw_t, f"{f} — raw data ({len(raw_t)} levels, {scope})", products, colours, thin_raw, base),
          "Every value, unbanded, after the data fixes (a numeric factor with more distinct values than "
          "raw_max_levels is shown in ~raw_bins fine bins instead). "
          f"Hollow markers have fewer than {thin_raw} claims. Click a product in the legend to show its own "
          "burning cost line.")
    tidy_note = (f"<p class='note'>The tidy banding Radar applies before Emblem — the same for every peril. "
                 f"Caps and floors at round numbers where the tails are thin (under {s['tail']:.1%} of exposure), "
                 f"and bands or levels with under {s['min_expo']:.1%} of exposure or {s['min_claims_tidy']} claims "
                 f"(all perils) merged. The rules are on the 'Banding rules' sheet of the spec. "
                 f"{'These are your override bands. ' if tidy.get('overridden') else ''}"
                 f"To change them, put a line like this in TIDY_OVERRIDES:</p><pre>{_esc(spec_line(f, tidy))}</pre>")
    b2_html = (tidy_note + HtmlReport._fig(plot_levels(
        tidy_t, f"{f} — suggested bands for modelling ({len(tidy_t)} levels, {scope})", products, colours,
        thin_raw, base)) + HtmlReport._blk("Table", small_table(tidy_t, thin_raw, views["tidy_base"]),
                                           open_=False))
    grp_note = (f"<p class='note'>A one-way guide to where grouping in Emblem may end up for {scope}: tidy bands "
                f"merged until each has at least {views['min_c']} claims (neighbours with the closest burning cost "
                f"first). Emblem's grouping, after fitting the other factors, decides. "
                f"Base level <b>{_esc(base or 'average')}</b> ({_esc(views['base_how'])}); every chart for this factor "
                f"is on that scale. To choose a different base, add <code>{_esc(json.dumps(f))}: "
                f"{_esc(json.dumps(base or ''))}</code> to BASE_LEVELS.</p>")
    b3_html = (grp_note + HtmlReport._fig(plot_levels(
        grp_t, f"{f} — indicative Emblem grouping ({len(grp_t)} groups, {scope})", products, colours,
        views["min_c"], base)) + HtmlReport._blk("Table", small_table(grp_t, views["min_c"], base), open_=False))
    blocks = [("1. Raw data", b1[1], b1[2]), ("2. Suggested bands for modelling (Radar)", b2_html, ""),
              (f"3. Indicative Emblem grouping ({scope})", b3_html, "")]
    per = cfg.get("PERIOD")
    if per and s.get("show_stability", True):
        cc, kc = ("_n_all", "_c_all") if peril == "ALL" else cfg["PERIL_COLS"][peril]
        lab = tabs["group"][1]
        d = pd.DataFrame({"lab": lab, "y": df[per], "e": df[cfg["EXPO"]], "c": df[kc]})
        g = d.groupby(["y", "lab"])[["e", "c"]].sum()
        yr = d.groupby("y")[["e", "c"]].sum()
        rel = (g.c / g.e) / (yr.c / yr.e).reindex(g.index.get_level_values(0)).to_numpy()
        rel = rel.unstack(0).reindex(list(grp_t.index))
        if base is not None and base in rel.index:
            b = rel.loc[base]
            rel = rel.div(b.where(b > 0), axis=1)
        st = plot_stability(rel, f"{f} — indicative groups by {per} ({scope}, each year vs its base)", base)
        st.update_layout(title=dict(y=0.98, yanchor="top"))
        blocks.append(("Stability by year (indicative groups)", fix_plot_area(st, list(rel.index), 260, 100), ""))
    rep.add_block_section(f, blocks, intro)
    # collapse the stability block by default
    if per and s.get("show_stability", True):
        rep.sections[-1] = rep.sections[-1].replace('<details class="blk" open><summary>Stability',
                                                    '<details class="blk"><summary>Stability')


# =============================================================================== Radar spec

def junk_condition(v, rule):
    parts = [f"{v} = {_num(x)}" for x in rule.get("values", [])]
    if "below" in rule:
        parts.append(f"{v} < {_num(rule['below'])}")
    if "above" in rule:
        parts.append(f"{v} > {_num(rule['above'])}")
    return " Or ".join(parts)


def _num(x):
    try:
        x = float(x)
        return str(int(x)) if x.is_integer() else repr(x)
    except (TypeError, ValueError):
        return f'"{x}"'


def radar_formula(f, cfg, na_code=None, junk_code=None):
    """Formula component text: N/A recode and/or numeric junk -> code (both mapped by Banding rules)."""
    rn = cfg.get("RADAR_NAMES", {})
    v = rn.get(f, f)
    rule = cfg.get("APPLIES_TO", {}).get(f)
    junk = cfg.get("JUNK", {}).get(f)
    if rule is None and junk is None:
        return None
    inner = v
    if junk is not None and junk_code is not None:
        inner = f"(If {junk_condition(v, junk)} then {_num(junk_code)} else {v})"
    if rule is not None:
        conds = []
        if "products" in rule:
            prod = rn.get(cfg["PRODUCT"], cfg["PRODUCT"])
            c = " And ".join(f'{prod} <> "{p}"' for p in rule["products"])
            conds.append(f"({c})" if len(rule["products"]) > 1 and "section" in rule else c)
        if "section" in rule:
            sc = cfg["SECTIONS"][rule["section"]]
            conds.append(f'{rn.get(sc, sc)} <> "Y"')
        code = f'"{na_code}"' if isinstance(na_code, str) else _num(na_code)
        return f"{v}Model = If {' Or '.join(conds)} then {code} else {inner};"
    return f"{v}Model = {inner[1:-1] if inner.startswith('(') else inner};"


def banding_rules(f, tidy, df_raw, df_eng, cfg, found, na_code, junk_code, x_values):
    """Rows for Radar's Banding grid: (Lower Op, Lower Bound, Upper Op, Upper Bound, Level name)."""
    rows = []
    if tidy["type"] == "numeric":
        codes = [c for c in (na_code, junk_code) if c is not None and not isinstance(c, str)]
        if na_code is not None:
            rows.append(("=", _tidy(na_code), "", "", "N/A"))
        if junk_code is not None:
            rows.append(("=", _tidy(junk_code), "", "", "Unknown"))
        rows.append(("NoValue", "", "", "", "Unknown"))
        e = tidy["edges"]
        names = numeric_labels(e, False, tidy["integer"], tidy.get("data_min"))
        if not e:
            rows.append((">", max(codes), "", "", names[0]) if codes else ("Default", "", "", "", names[0]))
        else:
            # lowest band: two-sided above the reserved codes so no rule overlaps another
            rows.append((">", max(codes), "<", e[0], names[0]) if codes else ("<", e[0], "", "", names[0]))
            for a, b, nm in zip(e, e[1:], names[1:-1]):
                rows.append((">=", a, "<", b, nm))
            rows.append((">=", e[-1], "", "", names[-1]))
        return rows
    if na_code is not None:
        rows.append(("=", "N/A", "", "", "N/A"))
    for ph in found.get(f, []):
        if ph != "":
            rows.append(("=", ph, "", "", "Unknown"))
    rows.append(("NoValue", "", "", "", "Unknown"))
    for raw_level, lvl in tidy["map"].items():
        rows.append(("=", raw_level, "", "", lvl))
    return rows


def write_spec(path, out_dir, cfg, factors, checks, tidy_specs, reserved, junk_codes, found, df_raw, df_eng,
               tidy_tables, version):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A78D6")

    def sheet(title, header, rows, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    sheet("Read me", ["Step", "Radar set-up"], [
        ("1", "Product: add the 'Product rules' to the Banding component (RPO / CPO / RE; NoValue and Default -> "
              "UNKNOWN as a catch-all). The Radar guide (34.3) warns Default rules can hide new levels after a "
              "refresh, so check the UNKNOWN count is zero after each refresh."),
        ("2", "Mappings: for any factor on the 'Mappings' sheet, read the mapping CSV and join on the raw value."),
        ("3", "Formula component (before Banding): add each formula on the 'Formulas' sheet. It creates <Factor>Model "
              "with reserved codes: the N/A code where the factor doesn't apply, and a junk code for placeholder / "
              "out-of-range values. Band <Factor>Model, not the raw field."),
        ("4", "Banding component: for each factor, enter the rows on 'Banding rules' (Lower Operator, Lower Bound, "
              "Upper Operator, Upper Bound, Level name). Rules don't overlap, so their order doesn't matter. Long "
              "categorical lists are also written as CSVs (tidy_map_<factor>.csv) to join instead of typing."),
        ("5", "Level names: enter them in the order on the 'Levels' sheet; tick Ordered for numeric factors; set the "
              "Base Level shown (used when Radar writes the Emblem data file)."),
        ("6", "In Emblem: fix the N/A parameter at 0 for factors with an N/A level. For variates, use the 'Variate x' "
              "values on the 'Levels' sheet as each level's x-value."),
    ], [8, 130], first=True)

    rows = []
    for c in checks:
        for lv, t in (c["flags"] or [("", "No issues found")]):
            rows.append((c["factor"], c["type"], c["distinct"], f"{c['missing_raw']:.1%}", f"{c['na_share']:.1%}",
                         f"{c['unknown']:.1%}", lv, t))
    sheet("Factor checks", ["Factor", "Type", "Distinct", "Missing (raw)", "N/A", "Unknown", "Level", "Finding"], rows,
          [22, 12, 10, 12, 8, 10, 8, 110])

    frows = []
    for f in factors:
        fm = radar_formula(f, cfg, reserved.get(f), junk_codes.get(f))
        if fm:
            frows.append((f, _rule_text(cfg["APPLIES_TO"][f]) if f in cfg.get("APPLIES_TO", {}) else "",
                          reserved.get(f, ""), junk_codes.get(f, ""), fm))
    sheet("Formulas", ["Factor", "Applies to", "N/A code", "Junk code", "Formula"],
          frows or [("(none)", "", "", "", "")], [22, 28, 10, 10, 110])

    brows, lrows = [], []
    for f in factors:
        t = tidy_specs[f]
        x = tidy_tables[f]
        rules = banding_rules(f, t, df_raw, df_eng, cfg, found, reserved.get(f), junk_codes.get(f), None)
        rtype = "Float" if pd.api.types.is_numeric_dtype(df_raw[f]) else "String"
        for r in rules:
            brows.append((f, rtype) + tuple(r))
        if t["type"] == "categorical" and len(t["map"]) > 50:
            pd.DataFrame({"raw": list(t["map"]), "level": list(t["map"].values())}).to_csv(
                out_dir / f"tidy_map_{f}_{version}.csv", index=False)
        tb = x["table"]
        for i, lvl in enumerate(tb.index, start=1):
            lrows.append((f, i, lvl, f"{tb.loc[lvl, 'expo_share']:.2%}",
                          x["xvals"].get(lvl, ""), "Y" if lvl == x["base"] else "",
                          "Y" if t["type"] == "numeric" else "N"))
    sheet("Banding rules", ["Factor", "Rules type", "Lower Operator", "Lower Bound", "Upper Operator", "Upper Bound",
                            "Level name"], brows, [22, 10, 14, 14, 14, 14, 22])
    sheet("Levels", ["Factor", "Order", "Level name", "Exposure %", "Variate x", "Base level", "Ordered"], lrows,
          [22, 7, 24, 11, 12, 10, 9])
    prows = [("=", p, "", "", p) for p in cfg["PRODUCTS"]] + [("NoValue", "", "", "", "UNKNOWN"),
                                                               ("Default", "", "", "", "UNKNOWN")]
    sheet("Product rules", ["Lower Operator", "Lower Bound", "Upper Operator", "Upper Bound", "Level name"], prows,
          [14, 14, 14, 14, 14])
    mrows = []
    for f, mp in cfg.get("MAPPINGS", {}).items():
        if isinstance(mp, str):
            tm = pd.read_csv(mp, dtype=str)
            mp = dict(zip(tm.iloc[:, 0], tm.iloc[:, 1]))
        mrows += [(f, k, v) for k, v in mp.items()]
    sheet("Mappings", ["Factor", "Raw value", "Mapped value"], mrows or [("(none)", "", "")], [22, 30, 30])
    wb.save(path)
    return path


def variate_x(df, f, lab, expo):
    """Exposure-weighted mean of the real values in each level (numeric factors)."""
    s = pd.to_numeric(df[f], errors="coerce")
    ok = ~lab.isin(SPECIAL) & s.notna()
    g = pd.DataFrame({"l": lab[ok], "xw": s[ok] * df.loc[ok, expo], "e": df.loc[ok, expo]}).groupby("l").sum()
    return {k: round(float(v), 2) for k, v in (g.xw / g.e).items()}


# =============================================================================== report summaries

def all_summary_html(cfg, prod_tab, unexpected, by_year, mix, cover_tab, cover_flags, overview):
    fmt = {"rows": ",.0f", "exposure": ",.0f", "expo_share": ".1%", "burning_cost": ",.2f"}
    fmt.update({c: ",.0f" for c in prod_tab.columns if c.startswith("claims_")})
    hdr = {"rows": "Rows", "exposure": "Exposure", "expo_share": "Exposure %", "burning_cost": "Burning cost"}
    hdr.update({c: c.replace("claims_", "Claims ") for c in prod_tab.columns if c.startswith("claims_")})
    html = "<h2>Products</h2>" + _table(prod_tab, fmt, hdr, index_name="Product")
    html += ("<p class='note'><b>Products outside the expected list:</b> " + ", ".join(_esc(u) for u in unexpected)
             + " — these fall to UNKNOWN in the Radar rules.</p>" if unexpected else
             "<p class='note'>No products outside " + ", ".join(cfg["PRODUCTS"])
             + " — the UNKNOWN catch-all is not triggered.</p>")
    if by_year is not None:
        by = by_year.copy()
        by.index = [str(_tidy(i)) for i in by.index]
        html += "<h3>Exposure by year</h3>" + _table(by, {c: ",.0f" for c in by.columns}, index_name=cfg["PERIOD"])
    if mix is not None:
        html += ("<h3>Section mix (share of each product's exposure)</h3>"
                 + _table(mix, {c: ".1%" for c in mix.columns}, index_name="Cover"))
    html += ("<h2>Peril cover check</h2><p class='note'>Actual ÷ expected claims by peril (expected at the overall "
             "frequency). Very low values with plenty of expected claims suggest the peril isn't covered there. "
             "Nothing is excluded automatically.</p>"
             + heat_table(cover_tab.drop(columns="exposure share"), "Product / cover", fmt=".2f"))
    html += ("<ul class='checks'>" + "".join(f"<li>{_esc(x)}</li>" for x in cover_flags) + "</ul>"
             if cover_flags else "<p class='note'>No product / cover combinations look uncovered.</p>")
    html += factor_table_html(overview, all_perils=True)
    return html


def factor_table_html(overview, all_perils=False):
    fmt = {"missing (raw)": ".1%", "N/A": ".1%", "Unknown": ".1%", "spread": ".2f"}
    hdr = {"type": "Type", "missing (raw)": "Missing (raw)", "actions": "Actions", "checks": "Checks",
           "bands": "Modelling bands", "groups": "Indicative groups", "spread": "Relativity spread",
           "base": "Base level"}
    return ("<h2>Factors</h2><p class='note'><b>Actions</b> need a decision before modelling (the factor's section "
            "gives the line to paste into the config); <b>checks</b> are worth a look. Relativity spread = highest ÷ "
            "lowest burning cost relativity across the indicative groups with enough claims — a rough one-way guide "
            "to how much a factor moves" + (" across all perils combined" if all_perils else "") + ".</p>"
            + _table(overview, fmt, hdr, index_name="Factor", links={i: _slug(i) for i in overview.index}))

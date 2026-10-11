"""
Unified review: data checks + three views per factor, one report for all perils and one per peril.

For each factor:
    1. Raw data                 every value (fine bins only for numeric factors with very many values);
                                N/A and Unknown as their own levels. No merging at all.
    2. Suggested bands (Radar)  a light-touch tidy-up for Emblem: one level per value wherever Emblem's
                                255-level limit allows, otherwise fine bins; only empty bands are merged;
                                categorical levels are only grouped (to OTHER) beyond the 255 limit.
                                The same for every peril - this is what Radar writes to the Emblem file.
    3. Indicative Emblem groups per peril, merging suggested bands until each has enough claims.
                                A one-way guide to where Emblem grouping may end up, not a decision.

Spikes ("point masses"): a value holding a lot of exposure compared with its neighbours (e.g. 0 years
trading, or a default sum insured) is always its own level and is never merged, in any view.

Every chart: exposure bars stacked by product; burning cost / frequency / severity relativity lines
(per-product burning cost lines can be switched on from the legend). All three charts for a factor use
the same scale: 1.00 = the base level (largest-exposure group, or your choice).

Level names used everywhere (and in the Radar spec):
    N/A      not applicable (product / section) - reserved code from the Formula component
    Unknown  missing where the factor applies, or a junk value
    OTHER    categorical levels grouped together (joins any OTHER level already in the data)
"""

import json
import math

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from banding import (C_REF, HtmlReport, _cell, _esc, _fmt, _is_whole, _slug, _table, _tidy, _weighted_quantiles,
                     heat_table, numeric_labels, plot_stability, quantile_edges)
from readiness import _json_rule, _rule_text

C_BC, C_FREQ, C_SEV, C_INK = "#1f1f1d", "#4a3aa7", "#e34948", "#1f1f1d"
SPECIAL = ["Unknown", "N/A"]
OTHER = "OTHER"
NOT_BASE = {"Unknown", "N/A", OTHER, "(other levels)"}


# =============================================================================== helpers

def _real(df, f, r):
    """Rows where the factor applies and has a value (not N/A, not missing)."""
    s = df[f]
    ok = s.notna()
    if r is not None:
        ok &= ~(s.astype("string") == str(r)).fillna(False) if isinstance(r, str) else ~(s == r).fillna(False)
    return ok


def _sig2(x, sf=2):
    """Round to sf significant figures: 41,372 -> 41,000 (2) or 41,400 (3); 0.0183 -> 0.018 exactly."""
    if x == 0 or not np.isfinite(x):
        return float(x)
    return float(f"{x:.{sf}g}")


def _between(a, b):
    """
    The shortest round number r with a < r <= b: a cut point between two neighbouring decimal values.
    E.g. between 0 and 0.01695 -> 0.01; between 0.65 and 0.6667 -> 0.66. A short number survives being
    copied into Excel and Radar exactly, whereas 0.6666666666666666 can round up past the value it should
    sit below and put that value in the wrong band.
    """
    if not (np.isfinite(a) and np.isfinite(b)) or b <= a:
        return b
    if b == 0:
        return 0.0
    for sf in range(1, 16):
        mag = 10 ** (math.floor(math.log10(abs(b))) - (sf - 1))
        r = float(f"{math.floor(b / mag) * mag:.{sf + 2}g}")
        if a < r <= b:
            return r
    return b


def _value_names(values):
    """Names for one-level-per-value levels, with enough decimals that no two values share a name."""
    vals = [float(v) for v in values]
    for dp in (4, 6, 8, 10, 12):
        names = [f"{v:,.0f}" if v.is_integer() else f"{v:,.{dp}f}".rstrip("0").rstrip(".") for v in vals]
        if len(set(names)) == len(names):
            return names
    return names


def _dedupe(seq):
    return list(dict.fromkeys(seq))


def find_points(v, w, point_min=0.02, point_ratio=3.0):
    """
    Spikes: values holding at least point_min of the exposure AND at least point_ratio times the average
    share of up to two neighbouring values on each side (so an even spread, like 60 years each with 1.7%,
    has no spikes, but 8% at 0 years trading does). Returns [(value, upper edge)], where the band
    [value, upper) holds only that value.
    """
    g = w.groupby(v).sum().sort_index()
    if len(g) < 2:
        return []
    share = (g / g.sum()).to_numpy()
    vals = g.index.to_numpy(float)
    whole = np.allclose(vals, np.round(vals))
    out = []
    for i, (x, sh) in enumerate(zip(vals, share)):
        if sh < point_min:
            continue
        nb = np.r_[share[max(0, i - 2):i], share[i + 1:i + 3]]
        if len(nb) and sh >= point_ratio * nb.mean():
            upper = x + 1 if whole else (_between(x, vals[i + 1]) if i + 1 < len(vals) else None)
            out.append((_tidy(x), None if upper is None else _tidy(upper)))
    return out


def _point_edges(points):
    """Edges that make each spike a band of its own: [value, upper)."""
    e = []
    for lo, up in points:
        e.append(lo)
        if up is not None:
            e.append(up)
    return e


def band_names(spec):
    """Level names for a numeric spec: per-value names, range labels, spikes named by their value."""
    e = spec["edges"]
    if spec.get("per_value"):
        names = _value_names(spec.get("values") or [spec["data_min"]] + list(e))
        if len(names) > 1:
            names[-1] = names[-1] + "+"
        return names
    names = numeric_labels(e, False, spec["integer"], spec.get("data_min"))
    for lo, up in spec.get("points", []):
        for i in range(len(names)):                   # band i holds [e[i-1], e[i]); band 0 everything below e[0]
            lo_i = e[i - 1] if i > 0 else spec.get("data_min")
            up_i = e[i] if i < len(e) else None
            if lo_i is not None and lo_i <= lo and up_i == up:   # the band that ends where the spike's does
                names[i] = _fmt(lo)
    seen = set()
    for i, nm in enumerate(names):                    # never two levels with the same name
        if nm in seen:
            names[i] = f"{nm} ({i})"
        seen.add(names[i])
    return names


def protected_edges(edges, points):
    """Edges either side of each spike: the edge its band starts at (the last at or below the value) and
    the next edge above it."""
    out = set()
    e = sorted(edges)
    for lo, _ in points:
        below = [x for x in e if x <= lo]
        if below:
            out.add(below[-1])
        above = [x for x in e if x > lo]
        if above:
            out.add(above[0])
    return out


def band_sums(df, f, edges, r, expo, claims, cost):
    """Exposure, claims and cost per band (applicable rows only), bands [e[i-1], e[i])."""
    ok = _real(df, f, r)
    v = df.loc[ok, f].astype(float).to_numpy()
    k = np.searchsorted(np.asarray(edges, float), v, side="right")
    K = len(edges) + 1
    sums = [np.bincount(k, weights=df.loc[ok, col].to_numpy(float), minlength=K) for col in (expo, claims, cost)]
    return sums


def merge_adjacent(edges, e, n, c, protected=(), min_e=0.0, min_n=0.0, max_bands=None, by_cost=True):
    """
    Merge neighbouring bands until none is below min_e exposure or min_n claims and there are at most
    max_bands. The band with the fewest claims goes first; it joins the neighbour with the closer burning
    cost (by_cost) or the smaller exposure. Edges in `protected` (spike boundaries) are never removed.
    """
    edges = list(edges)
    E, N, C = list(map(float, e)), list(map(float, n)), list(map(float, c))
    prot = set(protected)

    def rate(k):
        return math.log(max(C[k] / E[k], 1e-12)) if E[k] > 0 else float("nan")

    while True:
        K = len(E)
        thin = [k for k in range(K) if E[k] < min_e or N[k] < min_n]
        over = max_bands is not None and K > max_bands
        if not thin and not over:
            break
        key = lambda k: (N[k], E[k])
        cands = sorted(thin, key=key) + (sorted(set(range(K)) - set(thin), key=key) if over else [])
        merged = False
        for k in cands:
            opts = []
            if k > 0 and edges[k - 1] not in prot:
                opts.append(k - 1)                     # merge with the left neighbour
            if k < K - 1 and edges[k] not in prot:
                opts.append(k)                         # merge with the right neighbour
            if not opts:
                continue
            j = opts[0]
            if len(opts) == 2:
                left, right = k - 1, k + 1
                dl, dr = abs(rate(k) - rate(left)), abs(rate(k) - rate(right))
                if by_cost and np.isfinite(dl) and np.isfinite(dr) and not math.isclose(dl, dr):
                    j = opts[0] if dl < dr else opts[1]
                elif by_cost:
                    j = opts[0] if N[left] <= N[right] else opts[1]
                else:
                    j = opts[0] if E[left] <= E[right] else opts[1]
            E[j] += E[j + 1]; N[j] += N[j + 1]; C[j] += C[j + 1]
            del E[j + 1], N[j + 1], C[j + 1], edges[j]
            merged = True
            break
        if not merged:
            break
    return edges


# =============================================================================== level specs

def raw_spec(df, f, cat, r, expo, s):
    """
    As raw as possible: every categorical level (up to raw_max_cat_levels) and every numeric value
    (up to raw_max_levels distinct values). Only beyond that is a numeric factor binned, into ~raw_bins
    equal-exposure bins (edges to 3 significant figures), with spikes kept as their own level.
    """
    ok = _real(df, f, r)
    sub = df.loc[ok]
    if cat:
        g = sub.groupby(sub[f].astype(str))[expo].sum().sort_values(ascending=False)
        top = list(g.index[:s["raw_max_cat_levels"]])
        rest = "(other levels)"
        mp = {k: (k if k in top else rest) for k in g.index}
        return {"type": "categorical", "map": mp, "levels": _dedupe(top + ([rest] if len(g) > len(top) else [])),
                "other": rest, "reserved": r}
    vals = np.sort(sub[f].unique())
    if len(vals) <= s["raw_max_levels"]:
        return {"type": "values", "values": [_tidy(v) for v in vals], "reserved": r}
    whole = _is_whole(sub[f])
    pts = find_points(sub[f].astype(float), sub[expo].astype(float), s["point_min"], s["point_ratio"])
    rest = sub[~sub[f].isin([p for p, _ in pts])]          # bins share out the exposure left after spikes
    q = quantile_edges(rest if len(rest) else sub, f, s["raw_bins"], expo)
    edges = sorted({_sig2(x, 3) for x in q} | set(_point_edges(pts)))
    edges = [_tidy(x) for x in edges if vals[0] < x <= vals[-1]]
    return {"type": "numeric", "edges": edges, "integer": whole, "points": pts, "data_min": _tidy(vals[0]),
            "reserved": r}


def tidy_spec(df, f, cat, r, expo, claims, cost, s, override=None):
    """
    Light-touch suggested bands for Radar -> Emblem. Nothing is grouped unless it has to be:
      categorical  every level kept; only beyond max_levels (Emblem's 255 limit, less room for Unknown,
                   N/A and OTHER) do the smallest levels go to OTHER (min_expo can add a floor, default 0)
      numeric      one level per value when there are at most max_levels distinct values; otherwise
                   ~tidy_bins equal-exposure bins (2 significant figures) with the far tails (tail)
                   capped; spikes always their own level; only empty bands merged.
    override: {"edges": [...]} or {"levels": [...]} replaces the suggestion.
    """
    ok = _real(df, f, r)
    sub = df.loc[ok]
    tot = sub[expo].sum()
    if cat:
        g = sub.groupby(sub[f].astype(str))[expo].sum().sort_values(ascending=False)
        if override and "levels" in override:
            keep = [str(x) for x in override["levels"] if str(x) != OTHER]
        else:
            keep = [k for k in g.index if g[k] / tot >= s["min_expo"] or k == OTHER][: s["max_levels"]]
        mp = {k: (k if k in keep else OTHER) for k in g.index}
        grouped = [k for k, v in mp.items() if v == OTHER and k != OTHER]
        levels = _dedupe([k for k in g.index if k in keep and k != OTHER] + ([OTHER] if OTHER in mp.values() else []))
        return {"type": "categorical", "map": mp, "levels": levels, "other": OTHER, "reserved": r,
                "overridden": bool(override), "grouped": grouped, "data_has_other": OTHER in g.index}
    v = sub[f].astype(float)
    w = sub[expo].astype(float)
    vals = np.sort(v.unique())
    whole = _is_whole(v)
    pts = find_points(v, w, s["point_min"], s["point_ratio"])
    base = {"type": "numeric", "integer": whole, "data_min": _tidy(vals[0]) if len(vals) else None,
            "reserved": r, "points": pts, "overridden": bool(override)}
    if override and "edges" in override:
        return {**base, "edges": sorted(_tidy(x) for x in override["edges"]), "points": []}
    if len(vals) <= s["max_levels"]:                 # one level per value
        edges = vals[1:] if whole else [_between(a, b) for a, b in zip(vals, vals[1:])]
        return {**base, "edges": [_tidy(x) for x in edges], "values": [_tidy(x) for x in vals], "per_value": True}
    rest = sub[~v.isin([p for p, _ in pts])]               # bins share out the exposure left after spikes
    q = quantile_edges(rest if len(rest) else sub, f, s["tidy_bins"], expo)
    edges = {_sig2(x) for x in q}
    lo_q, hi_q = _weighted_quantiles(v.to_numpy(), w.to_numpy(), [s["tail"], 1 - s["tail"]])
    lo_c, hi_c = _sig2(lo_q), _sig2(hi_q)
    edges = {x for x in edges if lo_c <= x <= hi_c} | {lo_c, hi_c} | set(_point_edges(pts))
    edges = sorted(x for x in edges if vals[0] < x <= vals[-1])
    E, N, C = band_sums(df, f, edges, r, expo, claims, cost)
    edges = merge_adjacent(edges, E, N, C, protected=protected_edges(edges, pts), min_e=1e-12,
                           max_bands=s["max_levels"], by_cost=False)
    return {**base, "edges": [_tidy(x) for x in edges]}


def group_spec(df, f, tidy, expo, claims, cost, min_claims, max_groups, point_min):
    """Indicative Emblem grouping for one peril: merge suggested levels until each has min_claims."""
    r = tidy.get("reserved")
    if tidy["type"] == "numeric":
        pts = tidy.get("points") or []
        edges = list(tidy["edges"])
        E, N, C = band_sums(df, f, edges, r, expo, claims, cost)
        edges = merge_adjacent(edges, E, N, C, protected=protected_edges(edges, pts), min_n=min_claims,
                               max_bands=max_groups, by_cost=True)
        return {**tidy, "edges": [_tidy(x) for x in edges], "per_value": False, "points": pts}
    ok = _real(df, f, r)
    sub = df.loc[ok]
    lv = sub[f].astype(str).map(tidy["map"]).fillna(OTHER)
    agg = sub.groupby(lv)[[expo, claims]].sum()
    tot = agg[expo].sum()
    keep = [l for l in tidy["levels"] if l != OTHER and l in agg.index
            and (agg.loc[l, claims] >= min_claims or agg.loc[l, expo] / tot >= point_min)]
    mp = {k: (v if v in keep else OTHER) for k, v in tidy["map"].items()}
    levels = _dedupe(keep + ([OTHER] if OTHER in mp.values() else []))
    return {**tidy, "map": mp, "levels": levels}


def level_series(s, spec):
    """Label every row with its level under spec. Returns (labels, ordered level list)."""
    r = spec.get("reserved")
    if spec["type"] in ("values", "numeric"):
        na = (s == r).fillna(False) if r is not None and not isinstance(r, str) else pd.Series(False, index=s.index)
        v = pd.to_numeric(s, errors="coerce").where(~na)
        ok = v.notna().to_numpy()
        lab = np.full(len(s), "Unknown", dtype=object)
        if spec["type"] == "values":
            order = _value_names(spec["values"])
            names = dict(zip([float(x) for x in spec["values"]], order))
            lab[ok] = v[ok].astype(float).map(names).to_numpy()
        else:
            order = band_names(spec)
            idx = np.searchsorted(np.asarray(spec["edges"], float), v[ok].to_numpy(float), side="right")
            lab[ok] = np.asarray(order, dtype=object)[idx]
        lab[na.to_numpy()] = "N/A"
        lab = pd.Series(lab, index=s.index, dtype=object)
    else:
        st = s.astype("string")
        mapped = st.map(spec["map"])
        lab = mapped.where(mapped.notna(), spec.get("other", OTHER)).astype(object)
        lab[st.isna().to_numpy()] = "Unknown"
        lab[st.eq("N/A").fillna(False).to_numpy()] = "N/A"
        order = _dedupe(spec["levels"])
    present = set(lab.unique())
    return lab, [o for o in order if o in present and o not in SPECIAL] + [o for o in SPECIAL if o in present]


# =============================================================================== tables & charts

def level_table(df, lab, order, expo, claims, cost, prod, products):
    d = pd.DataFrame({"lab": lab, "e": df[expo], "n": df[claims], "c": df[cost], "p": df[prod].astype("string")})
    E, N, C = d.e.sum(), d.n.sum(), d.c.sum()
    g = d.groupby("lab")[["e", "n", "c"]].sum().reindex(order)
    t = pd.DataFrame({"exposure": g.e, "claims": g.n, "cost": g.c})
    t["expo_share"] = t.exposure / E
    t["freq_rel"] = (t.claims / t.exposure) / (N / E)
    t["sev_rel"] = (t.cost / t.claims.where(t.claims > 0)) / (C / N) if N else np.nan
    t["bc_rel"] = (t.cost / t.exposure) / (C / E) if C else np.nan
    pe = d.groupby(["lab", "p"])[["e", "c"]].sum()
    for p in products:
        if p in pe.index.get_level_values(1):
            x = pe.xs(p, level=1).reindex(order)
            t[f"expo_{p}"] = x.e.fillna(0)
            t[f"bc_{p}"] = (x.c / x.e) / (C / E) if C else np.nan
        else:
            t[f"expo_{p}"] = 0.0
            t[f"bc_{p}"] = np.nan
    return t


def other_share(t):
    return float(t.loc[OTHER, "expo_share"]) if OTHER in t.index else 0.0


def pick_base_level(t, requested=None):
    if requested is not None and str(requested) in map(str, t.index):
        return str(requested), "your choice"
    cand = t[~t.index.isin(NOT_BASE) & (t.claims > 0)]
    if not len(cand):
        return None, "none available"
    how = "largest exposure" if requested is None else f"'{requested}' not found, so largest exposure"
    return str(cand.exposure.idxmax()), how


def rescale(t, scale):
    t = t.copy()
    for c in [c for c in t.columns if c.endswith("_rel") or c.startswith("bc_")]:
        key = "bc_rel" if c.startswith("bc_") else c
        if scale.get(key) and np.isfinite(scale[key]) and scale[key] > 0:
            t[c] = t[c] / scale[key]
    return t


PLOT_H, PLOT_TOP, ASSUMED_W, TICK_PX, MAX_TICKS = 330, 112, 1100, 11, 30


def fix_plot_area(fig, labels, plot_h=PLOT_H, top=PLOT_TOP, max_ticks=MAX_TICKS):
    """
    Same plotting-area height on every chart: the space the level names need is added below it.
    Tick labels: all of them when they fit side by side; otherwise every k-th one, flat if they're short,
    else at 40 degrees, never more than max_ticks. N/A, Unknown and OTHER are always labelled (dropped to a
    second line if they would collide). Hovering still shows every level.
    """
    labels = [str(l) for l in labels]
    n = len(labels)
    char_w = TICK_PX * 0.55
    left, right = 72, 72
    plot_w = ASSUMED_W - left - right
    slot = plot_w / max(n, 1)
    longest = max((len(l) for l in labels), default=1)
    fit_flat = max(1, int(plot_w // (longest * char_w + 14)))
    if n <= fit_flat:
        angle, step = 0, 1
    elif fit_flat >= 12:
        angle, step = 0, math.ceil(n / min(fit_flat, max_ticks))
    else:
        angle, step = -40, math.ceil(n / min(max_ticks, max(1, int(plot_w // 24))))
    gap = 18 / slot if angle else None               # angled labels: minimum spacing in slots

    def span(i):
        w = (len(labels[i]) * char_w + 6) / slot if angle == 0 else gap
        return i + 0.5 - w / 2, i + 0.5 + w / 2

    lines = {0: [], 1: []}

    def free(i, line):
        a, b = span(i)
        return all(b <= c or a >= d for c, d in (span(j) for j in lines[line]))

    for i in (i for i, l in enumerate(labels) if l in NOT_BASE):          # special levels first
        lines[0 if free(i, 0) or angle else 1].append(i)
    for i in range(0, n, step):
        if labels[i] not in NOT_BASE and free(i, 0):
            lines[0].append(i)
    shown = sorted(lines[0] + lines[1])
    text = [("<br>" if i in lines[1] else "") + labels[i] for i in shown]
    longest_shown = max((len(labels[i]) for i in shown), default=1)
    sin, cos = abs(math.sin(math.radians(angle))), abs(math.cos(math.radians(angle)))
    bottom = int(longest_shown * char_w * sin + TICK_PX * (cos + 0.6) + 20 + (TICK_PX + 4) * bool(lines[1]))
    if angle and shown:                               # the first label leans left, past the y axis
        left = max(left, int(len(labels[shown[0]]) * char_w * cos - slot * (shown[0] + 0.5) + 12))
    fig.update_xaxes(tickangle=angle, automargin=False, tickfont=dict(size=TICK_PX),
                     tickmode="array", tickvals=[labels[i] for i in shown], ticktext=text)
    fig.update_yaxes(automargin=False)
    fig.update_layout(height=plot_h + top + bottom, margin=dict(t=top, b=bottom, l=left, r=right))
    return fig


def plot_levels(t, title, products, colours, min_claims, base=None, hollow=True):
    x = [str(i) for i in t.index]
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    for p in products:
        fig.add_trace(go.Bar(x=x, y=t[f"expo_{p}"], name=p, marker=dict(color=colours[p], opacity=0.38,
                                                                         line=dict(width=0)),
                             hovertemplate=f"{p} exposure %{{y:,.0f}}<extra></extra>"), secondary_y=False)
    thin = (t.claims.fillna(0) < min_claims).to_numpy() if hollow else np.zeros(len(x), bool)
    dense = len(x) > 60                      # hundreds of levels: lighter marks so the shape shows
    for c, name, col, width, dash, size in [("bc_rel", "Burning cost", C_BC, 1.5 if dense else 3, "solid", 4 if dense else 9),
                                            ("freq_rel", "Frequency", C_FREQ, 1 if dense else 1.5, "dot", 3 if dense else 6),
                                            ("sev_rel", "Severity", C_SEV, 1 if dense else 1.5, "dash", 3 if dense else 6)]:
        fig.add_trace(go.Scatter(
            x=x, y=t[c], name=name, mode="lines+markers", line=dict(color=col, width=width, dash=dash),
            marker=dict(size=size, color=col, symbol=np.where(thin, "circle-open", "circle")),
            customdata=np.c_[t.claims.fillna(0), t.expo_share.fillna(0) * 100],
            hovertemplate=f"{name} %{{y:.2f}}" + ("<br>Claims %{customdata[0]:,.0f} · exposure %{customdata[1]:.1f}%"
                                                  if c == "bc_rel" else "") + "<extra></extra>"), secondary_y=True)
    for p in products:
        if t[f"bc_{p}"].notna().any():
            fig.add_trace(go.Scatter(x=x, y=t[f"bc_{p}"], name=f"{p} burning cost", mode="lines+markers",
                                     visible="legendonly", line=dict(color=colours[p], width=1.8),
                                     marker=dict(size=6, color=colours[p]),
                                     hovertemplate=f"{p} burning cost %{{y:.2f}}<extra></extra>"), secondary_y=True)
    if base is not None and base in x:
        fig.add_trace(go.Scatter(x=[base], y=[1], mode="markers", name="Base level",
                                 marker=dict(symbol="diamond", size=13, color=C_INK, line=dict(color="white", width=2)),
                                 hoverinfo="skip"), secondary_y=True)
    fig.add_shape(type="line", xref="paper", x0=0, x1=1, yref="y2", y0=1, y1=1,
                  line=dict(color=C_REF, width=1, dash="dot"))
    fig.update_layout(title=dict(text=title, y=0.98, yanchor="top"), template="plotly_white", hovermode="x unified",
                      barmode="stack", bargap=0.12 if len(x) <= 60 else 0.02,
                      legend=dict(orientation="h", x=0, y=1.02, yanchor="bottom", traceorder="normal"))
    fig.update_xaxes(type="category")
    fig.update_yaxes(title_text="Exposure", secondary_y=False, showgrid=False, tickformat=",.0f")
    real = t.exposure[~t.index.isin(NOT_BASE)].max()
    big = t.exposure[t.index.isin(NOT_BASE)]
    if len(big) and real > 0 and big.max() > 2.5 * real:   # N/A / Unknown / OTHER would flatten every other bar
        cap = real * 1.3
        fig.update_yaxes(range=[0, cap], secondary_y=False)
        for lvl, e in big[big > cap].items():
            fig.add_annotation(x=str(lvl), y=cap, yref="y", text=f"{e:,.0f} ↑", showarrow=False, yanchor="top",
                               textangle=-90 if len(x) > 40 else 0, font=dict(size=10, color="#5f5e5a"))
    fig.update_yaxes(title_text="Relativity (1 = base)" if base else "Relativity (1 = average)", secondary_y=True,
                     rangemode="tozero", tickmode="auto", nticks=6, tickformat=".1f", showgrid=True)
    return fix_plot_area(fig, x)


_TFMT = {"exposure": ",.0f", "expo_share": ".1%", "claims": ",.0f", "freq_rel": ".2f", "sev_rel": ".2f", "bc_rel": ".2f"}
_THDR = {"exposure": "Exposure", "expo_share": "Exposure %", "claims": "Claims", "freq_rel": "Frequency rel.",
         "sev_rel": "Severity rel.", "bc_rel": "Burning cost rel."}


def small_table(t, min_claims, base=None):
    tt = t[list(_TFMT)].copy()
    if base is not None:
        tt.index = [f"{i}  (base)" if str(i) == base else i for i in tt.index]
    return _table(tt, _TFMT, _THDR, ["thin" if n < min_claims else "" for n in t.claims.fillna(0)], index_name="Level")


# =============================================================================== per-factor build

def spec_line(f, spec):
    if spec["type"] == "numeric":
        return f'"{f}": {json.dumps({"edges": spec["edges"]})},'
    return f'"{f}": {json.dumps({"levels": [l for l in spec["levels"] if l != OTHER]})},'


def factor_views(df, f, cat, tidy, raw, peril, cfg, s):
    """Tables, base and figures for one factor and one peril ('ALL' = all perils combined)."""
    expo, prod, products = cfg["EXPO"], cfg["PRODUCT"], cfg["PRODUCTS"]
    cc, kc = ("_n_all", "_c_all") if peril == "ALL" else cfg["PERIL_COLS"][peril]
    min_c = s["min_claims_group"].get(peril, s["min_claims_group"]["default"])
    grp = group_spec(df, f, tidy, expo, cc, kc, min_c, s["max_groups"], s["point_min"])
    tabs = {}
    for k, sp in (("raw", raw), ("tidy", tidy), ("group", grp)):
        lab, order = level_series(df[f], sp)
        tabs[k] = (level_table(df, lab, order, expo, cc, kc, prod, products), lab)
    req = cfg.get("BASE_LEVELS", {}).get(f)
    tidy_base = None
    if req is not None and str(req) in map(str, tabs["tidy"][0].index):
        tidy_base = str(req)                          # chosen base is a suggested band: base group contains it
        base = str(pd.crosstab(tabs["tidy"][1], tabs["group"][1]).loc[tidy_base].idxmax())
        how = "your choice"
    else:
        base, how = pick_base_level(tabs["group"][0], req)
        if base is not None:                        # Radar base: largest suggested band inside the base group
            inside = tabs["tidy"][1][tabs["group"][1] == base]
            ex = df.loc[inside.index, expo].groupby(inside).sum()
            ex = ex[~ex.index.isin(NOT_BASE)]
            tidy_base = str(ex.idxmax()) if len(ex) else None
    if base is not None:
        row = tabs["group"][0].loc[base]
        scale = {"bc_rel": row.bc_rel, "freq_rel": row.freq_rel, "sev_rel": row.sev_rel}
        tabs = {k: (rescale(t, scale), lab) for k, (t, lab) in tabs.items()}
    gt = tabs["group"][0]
    cred = gt[(gt.claims >= min_c) & ~gt.index.isin(SPECIAL)]
    spread = (cred.bc_rel.max() / cred.bc_rel.min()) if len(cred) > 1 and cred.bc_rel.min() > 0 else np.nan
    return dict(group=grp, tabs=tabs, base=base, base_how=how, tidy_base=tidy_base, min_c=min_c, spread=spread)


def factor_flags(tidy, tidy_t):
    """Extra findings from the suggested bands (level limit, OTHER share, spikes kept, OTHER in the data)."""
    out = []
    n = len(tidy_t)
    if n > 255:
        out.append(("action", f"{n} modelling levels — over Emblem's 255 limit. Group further (MAPPINGS to a parent)."))
    oth = other_share(tidy_t)
    if tidy["type"] == "categorical" and tidy.get("grouped"):
        out.append(("action" if oth >= 0.2 else "info",
                    f"{len(tidy['grouped'])} small levels grouped to OTHER to stay within Emblem's 255-level limit "
                    f"(OTHER = {oth:.0%} of exposure). A parent grouping via MAPPINGS keeps more detail."))
    if tidy.get("data_has_other"):
        out.append(("info", "The data already has a level called OTHER; any grouped levels join it (no duplicate)."))
    pts = tidy.get("points") or []
    if pts:
        out.append(("info", "Kept as their own level (spikes in exposure): " + ", ".join(_fmt(p) for p, _ in pts)))
    return out


def add_factor_section(rep, df, f, cat, chk, eng_log, raw, tidy, views, peril, cfg, colours, s, show, formula=None):
    """show: dict of report switches (factor_details, raw_view, tidy_view, group_view, stability, tables)."""
    products = cfg["PRODUCTS"]
    tabs = views["tabs"]
    lvl_word = {"action": "Action", "check": "Check", "info": "Info"}
    raw_t, tidy_t, grp_t = tabs["raw"][0], tabs["tidy"][0], tabs["group"][0]
    grp = views["group"]
    base = views["base"]
    scope = "all perils" if peril == "ALL" else peril
    intro = ""
    if show.get("factor_details", True):
        tiles = [("Type", chk["type"]), ("Distinct", f"{chk['distinct']:,}"),
                 ("Missing (raw)", f"{chk['missing_raw']:.1%}"), ("N/A", f"{chk['na_share']:.1%}"),
                 ("Unknown", f"{chk['unknown']:.1%}"), ("Modelling levels", f"{len(tidy_t)}"),
                 ("Indicative groups", f"{len(grp_t)}"), ("Base level", base or "average")]
        intro = ("<div class='stats'>" + "".join(f"<div class='stat'><span>{k}</span><b>{_esc(v)}</b></div>"
                                                  for k, v in tiles) + "</div>")
        flags = chk["flags"] + factor_flags(tidy, tidy_t)
        intro += "<ul class='checks'>" + ("".join(f"<li><b>{lvl_word[lv]}:</b> {_esc(t)}</li>" for lv, t in flags)
                                          or "<li>No data issues found.</li>") + "</ul>"
        if eng_log.get(f):
            intro += "<h3>Engineering applied</h3><ul class='checks'>" + "".join(
                f"<li>{_esc(t)}</li>" for t in eng_log[f]) + "</ul>"
        for p in chk.get("proposals", []):
            intro += ("<p class='note'>To apply the proposed recode, add this line to APPLIES_TO and run again:</p>"
                      f"<pre>{_esc(chr(34) + f + chr(34) + ': ' + _json_rule(p['rule']) + ',')}</pre>")
        if formula:
            intro += f"<p class='note'>Radar Formula component (before Banding):</p><pre>{_esc(formula)}</pre>"
        if tidy.get("reserved") is not None:
            intro += ("<p class='note'>In Emblem, fix the N/A level's parameter at 0 (relativity 1.00) with Fixing "
                      "Factors, so the product / section relativity carries those policies.</p>")

    blocks = []
    if show.get("raw_view", True):
        blocks.append(("1. Raw data", plot_levels(raw_t, f"{f} — raw data ({len(raw_t)} levels, {scope})", products,
                                                  colours, 0, base, hollow=False),
                       "Every value, as it is after the data fixes — nothing merged, levels with no claims included. "
                       "Only a numeric factor with more distinct values than raw_max_levels is shown in fine bins. "
                       "Hover over a level to see it exactly. Click a product in the legend to show its own "
                       "burning cost line."))
    if show.get("tidy_view", True):
        if tidy["type"] == "numeric" and tidy.get("per_value"):
            how = "one level per value — nothing grouped"
        elif tidy["type"] == "numeric" and not tidy.get("overridden"):
            how = (f"{len(tidy_t)} fine bands (too many distinct values for one level each within Emblem's limit); "
                   f"only the outer {s['tail']:.1%} tails capped and empty bands merged")
        elif tidy["type"] == "categorical" and not tidy.get("grouped"):
            how = "every level kept — nothing grouped"
        else:
            how = "your override" if tidy.get("overridden") else "smallest levels grouped to OTHER for Emblem's limit"
        line = spec_line(f, tidy)
        tidy_note = (f"<p class='note'>The light-touch tidy-up Radar applies before Emblem, the same for every peril: "
                     f"<b>{_esc(how)}</b>. Spikes are always their own level. The rules are on the 'Banding rules' "
                     f"sheet of the spec." + (f" To change them, put a line like this in TIDY_OVERRIDES:</p>"
                                               f"<pre>{_esc(line)}</pre>" if len(tidy_t) <= 60 else
                                               " To change them, add an entry to TIDY_OVERRIDES "
                                               "(<code>{\"edges\": [...]}</code> or <code>{\"levels\": [...]}"
                                               "</code>).</p>"))
        b2 = tidy_note + HtmlReport._fig(plot_levels(
            tidy_t, f"{f} — suggested levels for modelling ({len(tidy_t)} levels, {scope})", products, colours,
            0, base, hollow=False))
        if show.get("tables", True):
            b2 += HtmlReport._blk("Table", small_table(tidy_t, 0, views["tidy_base"]), open_=False)
        blocks.append(("2. Suggested levels for modelling (Radar)", b2, ""))
    if show.get("group_view", True):
        grp_note = (f"<p class='note'>A one-way guide to where grouping in Emblem may end up for {scope}: suggested "
                    f"levels merged until each has at least {views['min_c']} claims (neighbours with the closest "
                    f"burning cost first; spikes and large categorical levels never merged). Emblem's grouping, "
                    f"after fitting the other factors, decides. Base level <b>{_esc(base or 'average')}</b> "
                    f"({_esc(views['base_how'])}); every chart for this factor is on that scale. To choose a "
                    f"different base, add <code>{_esc(json.dumps(f))}: {_esc(json.dumps(base or ''))}</code> to "
                    f"BASE_LEVELS. Hollow markers have fewer than {views['min_c']} claims.</p>")
        if grp["type"] == "categorical" and other_share(grp_t) >= 0.5:
            grp_note += (f"<p class='note'><b>Most levels are too small to stand alone for {_esc(scope)}</b> "
                         f"(OTHER = {other_share(grp_t):.0%} of exposure). Emblem will need a parent grouping "
                         f"(e.g. a MAPPINGS to a broader group) or levels grouped by similar relativity.</p>")
        b3 = grp_note + HtmlReport._fig(plot_levels(
            grp_t, f"{f} — indicative Emblem grouping ({len(grp_t)} group{'' if len(grp_t) == 1 else 's'}, {scope})", products, colours,
            views["min_c"], base))
        if show.get("tables", True):
            b3 += HtmlReport._blk("Table", small_table(grp_t, views["min_c"], base), open_=False)
        blocks.append((f"3. Indicative Emblem grouping ({scope})", b3, ""))
    per = cfg.get("PERIOD")
    if per and show.get("stability", True):
        cc, kc = ("_n_all", "_c_all") if peril == "ALL" else cfg["PERIL_COLS"][peril]
        lab = tabs["group"][1]
        d = pd.DataFrame({"lab": lab, "y": df[per], "e": df[cfg["EXPO"]], "c": df[kc]})
        g = d.groupby(["y", "lab"])[["e", "c"]].sum()
        yr = d.groupby("y")[["e", "c"]].sum()
        rel = (g.c / g.e) / (yr.c / yr.e).reindex(g.index.get_level_values(0)).to_numpy()
        rel = rel.unstack(0).reindex(list(grp_t.index))
        if base is not None and base in rel.index:
            b = rel.loc[base]
            rel = rel.div(b.where(b > 0), axis=1)
        st = plot_stability(rel, f"{f} — indicative groups by {per} ({scope}, each year vs its base)", base)
        st.update_layout(title=dict(y=0.98, yanchor="top"))
        blocks.append(("Stability by year (indicative groups)", fix_plot_area(st, list(rel.index), 260, 100), ""))
    if not blocks and not intro:
        return False
    rep.add_block_section(f, blocks, intro)
    if per and show.get("stability", True):        # collapse the stability block by default
        rep.sections[-1] = rep.sections[-1].replace('<details class="blk" open><summary>Stability',
                                                    '<details class="blk"><summary>Stability')
    return True


# =============================================================================== Radar spec

def junk_condition(v, rule):
    parts = [f"{v} = {_num(x)}" for x in rule.get("values", [])]
    if "below" in rule:
        parts.append(f"{v} < {_num(rule['below'])}")
    if "above" in rule:
        parts.append(f"{v} > {_num(rule['above'])}")
    return " Or ".join(parts)


def _num(x):
    try:
        x = float(x)
        return str(int(x)) if x.is_integer() else repr(x)
    except (TypeError, ValueError):
        return f'"{x}"'


def radar_formula(f, cfg, na_code=None, junk_code=None):
    """Formula component text: N/A recode and/or numeric junk -> code (both mapped by Banding rules)."""
    rn = cfg.get("RADAR_NAMES", {})
    v = rn.get(f, f)
    rule = cfg.get("APPLIES_TO", {}).get(f)
    junk = cfg.get("JUNK", {}).get(f)
    if rule is None and junk is None:
        return None
    inner = v
    if junk is not None and junk_code is not None:
        inner = f"(If {junk_condition(v, junk)} then {_num(junk_code)} else {v})"
    if rule is not None:
        conds = []
        if "products" in rule:
            prod = rn.get(cfg["PRODUCT"], cfg["PRODUCT"])
            c = " And ".join(f'{prod} <> "{p}"' for p in rule["products"])
            conds.append(f"({c})" if len(rule["products"]) > 1 and "section" in rule else c)
        if "section" in rule:
            sc = cfg["SECTIONS"][rule["section"]]
            conds.append(f'{rn.get(sc, sc)} <> "Y"')
        code = f'"{na_code}"' if isinstance(na_code, str) else _num(na_code)
        return f"{v}Model = If {' Or '.join(conds)} then {code} else {inner};"
    return f"{v}Model = {inner[1:-1] if inner.startswith('(') else inner};"


def banding_rules(f, tidy, df_raw, df_eng, cfg, found, na_code, junk_code):
    """Rows for Radar's Banding grid: (Lower Op, Lower Bound, Upper Op, Upper Bound, Level name)."""
    rows = []
    if tidy["type"] == "numeric":
        codes = [c for c in (na_code, junk_code) if c is not None and not isinstance(c, str)]
        if na_code is not None:
            rows.append(("=", _tidy(na_code), "", "", "N/A"))
        if junk_code is not None:
            rows.append(("=", _tidy(junk_code), "", "", "Unknown"))
        rows.append(("NoValue", "", "", "", "Unknown"))
        e = tidy["edges"]
        names = band_names(tidy)
        if not e:
            rows.append((">", max(codes), "", "", names[0]) if codes else ("Default", "", "", "", names[0]))
        else:
            # lowest band: two-sided above the reserved codes so no rule overlaps another
            rows.append((">", max(codes), "<", e[0], names[0]) if codes else ("<", e[0], "", "", names[0]))
            for a, b, nm in zip(e, e[1:], names[1:-1]):
                rows.append((">=", a, "<", b, nm))
            rows.append((">=", e[-1], "", "", names[-1]))
        return rows
    if na_code is not None:
        rows.append(("=", "N/A", "", "", "N/A"))
    for ph in found.get(f, []):
        if ph != "":
            rows.append(("=", ph, "", "", "Unknown"))
    rows.append(("NoValue", "", "", "", "Unknown"))
    for raw_level, lvl in tidy["map"].items():
        rows.append(("=", raw_level, "", "", lvl))
    return rows


def write_spec(path, out_dir, cfg, factors, checks, tidy_specs, reserved, junk_codes, found, df_raw, df_eng,
               tidy_tables, version):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A78D6")

    def sheet(title, header, rows, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    sheet("Read me", ["Step", "Radar set-up"], [
        ("1", "Product: add the 'Product rules' to the Banding component (RPO / CPO / RE; NoValue and Default -> "
              "UNKNOWN as a catch-all). The Radar guide (34.3) warns Default rules can hide new levels after a "
              "refresh, so check the UNKNOWN count is zero after each refresh."),
        ("2", "Mappings: for any factor on the 'Mappings' sheet, read the mapping CSV and join on the raw value."),
        ("3", "Formula component (before Banding): add each formula on the 'Formulas' sheet. It creates <Factor>Model "
              "with reserved codes: the N/A code where the factor doesn't apply, and a junk code for placeholder / "
              "out-of-range values. Band <Factor>Model, not the raw field."),
        ("4", "Banding component: for each factor, enter the rows on 'Banding rules'. Rules don't overlap, so their "
              "order doesn't matter. Factors marked 'one level per value' on the Levels sheet can also be generated "
              "with Radar's Generate > Unique option. Long categorical lists are also written as CSVs "
              "(tidy_map_<factor>.csv) to join instead of typing."),
        ("5", "Level names: enter them in the order on the 'Levels' sheet; tick Ordered for numeric factors; set the "
              "Base Level shown (used when Radar writes the Emblem data file)."),
        ("6", "In Emblem: fix the N/A parameter at 0 for factors with an N/A level. For variates, use the 'Variate x' "
              "values on the 'Levels' sheet as each level's x-value."),
    ], [8, 130], first=True)

    rows = []
    for c in checks:
        extra = factor_flags(tidy_specs[c["factor"]], tidy_tables[c["factor"]]["table"])
        for lv, t in (c["flags"] + extra or [("", "No issues found")]):
            rows.append((c["factor"], c["type"], c["distinct"], f"{c['missing_raw']:.1%}", f"{c['na_share']:.1%}",
                         f"{c['unknown']:.1%}", lv, t))
    sheet("Factor checks", ["Factor", "Type", "Distinct", "Missing (raw)", "N/A", "Unknown", "Level", "Finding"], rows,
          [22, 12, 10, 12, 8, 10, 8, 110])

    frows = []
    for f in factors:
        fm = radar_formula(f, cfg, reserved.get(f), junk_codes.get(f))
        if fm:
            frows.append((f, _rule_text(cfg["APPLIES_TO"][f]) if f in cfg.get("APPLIES_TO", {}) else "",
                          reserved.get(f, ""), junk_codes.get(f, ""), fm))
    sheet("Formulas", ["Factor", "Applies to", "N/A code", "Junk code", "Formula"],
          frows or [("(none)", "", "", "", "")], [22, 28, 10, 10, 110])

    brows, lrows = [], []
    for f in factors:
        t = tidy_specs[f]
        x = tidy_tables[f]
        rules = banding_rules(f, t, df_raw, df_eng, cfg, found, reserved.get(f), junk_codes.get(f))
        rtype = "Float" if pd.api.types.is_numeric_dtype(df_raw[f]) else "String"
        for r in rules:
            brows.append((f, rtype) + tuple(r))
        if t["type"] == "categorical" and len(t["map"]) > 50:
            pd.DataFrame({"raw": list(t["map"]), "level": list(t["map"].values())}).to_csv(
                out_dir / f"tidy_map_{f}_{version}.csv", index=False)
        tb = x["table"]
        kind = ("one level per value" if t.get("per_value") else
                "every level" if t["type"] == "categorical" and not t.get("grouped") else "")
        for i, lvl in enumerate(tb.index, start=1):
            lrows.append((f, i, lvl, f"{tb.loc[lvl, 'expo_share']:.2%}",
                          x["xvals"].get(lvl, ""), "Y" if lvl == x["base"] else "",
                          "Y" if t["type"] == "numeric" else "N", kind if i == 1 else ""))
    sheet("Banding rules", ["Factor", "Rules type", "Lower Operator", "Lower Bound", "Upper Operator", "Upper Bound",
                            "Level name"], brows, [22, 10, 14, 14, 14, 14, 22])
    sheet("Levels", ["Factor", "Order", "Level name", "Exposure %", "Variate x", "Base level", "Ordered", "Note"],
          lrows, [22, 7, 24, 11, 12, 10, 9, 22])
    prows = [("=", p, "", "", p) for p in cfg["PRODUCTS"]] + [("NoValue", "", "", "", "UNKNOWN"),
                                                               ("Default", "", "", "", "UNKNOWN")]
    sheet("Product rules", ["Lower Operator", "Lower Bound", "Upper Operator", "Upper Bound", "Level name"], prows,
          [14, 14, 14, 14, 14])
    mrows = []
    for f, mp in cfg.get("MAPPINGS", {}).items():
        if isinstance(mp, str):
            tm = pd.read_csv(mp, dtype=str)
            mp = dict(zip(tm.iloc[:, 0], tm.iloc[:, 1]))
        mrows += [(f, k, v) for k, v in mp.items()]
    sheet("Mappings", ["Factor", "Raw value", "Mapped value"], mrows or [("(none)", "", "")], [22, 30, 30])
    wb.save(path)
    return path


def variate_x(df, f, lab, expo):
    """Exposure-weighted mean of the real values in each level (numeric factors)."""
    s = pd.to_numeric(df[f], errors="coerce")
    ok = ~lab.isin(SPECIAL) & s.notna()
    g = pd.DataFrame({"l": lab[ok], "xw": s[ok] * df.loc[ok, expo], "e": df.loc[ok, expo]}).groupby("l").sum()
    return {k: round(float(v), 2) for k, v in (g.xw / g.e).items()}


# =============================================================================== report summaries

def products_html(cfg, prod_tab, unexpected, by_year, mix):
    fmt = {"rows": ",.0f", "exposure": ",.0f", "expo_share": ".1%", "burning_cost": ",.2f"}
    fmt.update({c: ",.0f" for c in prod_tab.columns if c.startswith("claims_")})
    hdr = {"rows": "Rows", "exposure": "Exposure", "expo_share": "Exposure %", "burning_cost": "Burning cost"}
    hdr.update({c: c.replace("claims_", "Claims ") for c in prod_tab.columns if c.startswith("claims_")})
    html = "<h2>Products</h2>" + _table(prod_tab, fmt, hdr, index_name="Product")
    html += ("<p class='note'><b>Products outside the expected list:</b> " + ", ".join(_esc(u) for u in unexpected)
             + " — these fall to UNKNOWN in the Radar rules.</p>" if unexpected else
             "<p class='note'>No products outside " + ", ".join(cfg["PRODUCTS"])
             + " — the UNKNOWN catch-all is not triggered.</p>")
    if by_year is not None:
        by = by_year.copy()
        by.index = [str(_tidy(i)) for i in by.index]
        html += "<h3>Exposure by year</h3>" + _table(by, {c: ",.0f" for c in by.columns}, index_name=cfg["PERIOD"])
    if mix is not None:
        html += ("<h3>Section mix (share of each product's exposure)</h3>"
                 + _table(mix, {c: ".1%" for c in mix.columns}, index_name="Cover"))
    return html


def cover_html(cover_tab, cover_flags):
    html = ("<h2>Peril cover check</h2><p class='note'>Actual ÷ expected claims by peril (expected at the overall "
            "frequency). Very low values with plenty of expected claims suggest the peril isn't covered there. "
            "Nothing is excluded automatically.</p>"
            + heat_table(cover_tab.drop(columns="exposure share"), "Product / cover", fmt=".2f"))
    return html + ("<ul class='checks'>" + "".join(f"<li>{_esc(x)}</li>" for x in cover_flags) + "</ul>"
                   if cover_flags else "<p class='note'>No product / cover combinations look uncovered.</p>")


def factor_table_html(overview, all_perils=False):
    fmt = {"missing (raw)": ".1%", "N/A": ".1%", "Unknown": ".1%", "spread": ".2f", "importance": ".1f"}
    hdr = {"type": "Type", "missing (raw)": "Missing (raw)", "actions": "Actions", "checks": "Checks",
           "bands": "Modelling levels", "groups": "Indicative groups", "spread": "Relativity spread",
           "base": "Base level", "importance": "Importance", "imp_rank": "Rank"}
    return ("<h2>Factors</h2><p class='note'><b>Actions</b> need a decision before modelling (the factor's section "
            "gives the line to paste into the config); <b>checks</b> are worth a look. Relativity spread = highest ÷ "
            "lowest burning cost relativity across the indicative groups with enough claims — a rough one-way guide "
            "to how much a factor moves" + (" across all perils combined" if all_perils else "") + ". Importance "
            "(% of the top factor) and rank come from the gradient boosting model (see Feature importance) and "
            "allow for the other factors; '(noise)' = no better than random numbers.</p>"
            + _table(overview, fmt, hdr, index_name="Factor", links={i: _slug(i) for i in overview.index}))
